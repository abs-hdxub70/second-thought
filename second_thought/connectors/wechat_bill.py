"""WeChat Pay (微信支付) bill export parser, CSV or XLSX.

Users request the bill in WeChat (我 > 服务 > 钱包 > 账单 > 下载账单). Depending on the app
version it arrives as CSV (UTF-8) or XLSX. Typical columns:
交易时间, 交易类型, 交易对方, 商品, 收/支, 金额(元), 支付方式, 当前状态, 交易单号, 商户单号, 备注
Amounts carry a ¥ sign; neutral rows use "/" in the 收/支 column.
"""

from __future__ import annotations

import io
from pathlib import Path

from ..engine.categories import enrich
from ..models import EXPENSE, INCOME, NEUTRAL, Transaction
from .base import find_table, parse_amount, parse_datetime, read_text, stable_id

REQUIRED = ("交易时间", "交易类型", "收/支")
SKIP_STATUS = ("已全额退款", "对方已退还", "交易关闭", "支付失败")


def _direction(value: str, txn_type: str) -> str:
    if "支出" in value:
        return EXPENSE
    if "收入" in value:
        return INCOME
    return NEUTRAL


def _xlsx_rows(data: bytes) -> list[list[str]]:
    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    ws = wb.active
    rows: list[list[str]] = []
    started = False
    for raw in ws.iter_rows(values_only=True):
        cells = ["" if v is None else str(v).strip() for v in raw]
        if not started and all(any(col in c for c in cells) for col in REQUIRED):
            started = True
        if started and any(cells):
            rows.append(cells)
    if not rows:
        raise ValueError("Header row not found in WeChat XLSX export.")
    return rows


class WeChatBillConnector:
    name = "wechat"
    region = "CN"
    status = "real-parser"

    def load(self, source: str | Path | bytes) -> list[Transaction]:
        data = source if isinstance(source, bytes) else Path(source).read_bytes()
        if data[:2] == b"PK":  # zip container = xlsx
            rows = _xlsx_rows(data)
        else:
            rows = find_table(read_text(data), REQUIRED)
        header, body = rows[0], rows[1:]
        idx = {name.strip(): i for i, name in enumerate(header)}
        amount_col = next((h for h in idx if h.startswith("金额")), "金额(元)")

        def col(row: list[str], name: str) -> str:
            i = idx.get(name)
            return row[i] if i is not None and i < len(row) else ""

        out: list[Transaction] = []
        for row in body:
            when = col(row, "交易时间")
            if not when or not when[:4].isdigit():
                continue
            status = col(row, "当前状态")
            if any(s in status for s in SKIP_STATUS):
                continue
            txn_type = col(row, "交易类型")
            ref = col(row, "交易单号")
            txn = Transaction(
                id="wx-" + (ref or stable_id(when, col(row, "交易对方"), col(row, amount_col))),
                source=self.name,
                timestamp=parse_datetime(when),
                amount=parse_amount(col(row, amount_col)),
                currency="CNY",
                direction=_direction(col(row, "收/支"), txn_type),
                merchant=col(row, "交易对方"),
                description=col(row, "商品"),
                raw_category=txn_type,
                category="other",
                payment_method=col(row, "支付方式"),
                status=status,
                raw={k: col(row, k) for k in idx},
            )
            if txn_type in ("微信红包", "转账", "零钱提现", "零钱充值") and txn.direction == NEUTRAL:
                txn.category = "transfer"
            out.append(enrich(txn))
        return out
